# -*- coding: utf-8 -*-
"""分类轴 + Excel 知识库导出（2026-09-07 面试官口径新增）回归测试。"""
import os
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from food_categories import classify_text, is_drink, keyword_to_category
from export_knowledge_base import HEADERS, build_rows, collect_food_cards, write_xlsx


class CategoryAxisTests(unittest.TestCase):
    def test_salmon_bowl_is_meal_japanese(self):
        cat = classify_text("三文鱼波奇饭", "减脂期也能吃的满足感晚餐 1比1复刻")
        self.assertEqual(cat["meal"], "正餐")
        self.assertEqual(cat["sub"], "日料")
        self.assertFalse(cat["excluded_drink"])

    def test_milk_tea_is_drink_and_excluded(self):
        cat = classify_text("奶茶新品", "黑糖波霸 芋圆 生椰")
        self.assertTrue(cat["excluded_drink"])
        self.assertEqual(cat["drink_kind"], "奶茶茶饮")

    def test_keyword_direct_mapping(self):
        self.assertTrue(is_drink("奶茶新品"))
        self.assertFalse(is_drink("三文鱼波奇饭"))
        self.assertEqual(keyword_to_category("三文鱼波奇饭")["sub"], "日料")

    def test_random_word_not_forced_into_category(self):
        cat = classify_text("随便搜的随机词")
        self.assertFalse(cat["excluded_drink"])
        self.assertEqual(cat["meal"], "")


class KnowledgeBaseExportTests(unittest.TestCase):
    def test_headers_match_interviewer_columns(self):
        expected_first = ["渠道来源", "一级品类", "二级品类", "美食话题", "标题"]
        self.assertEqual(HEADERS[:5], expected_first)
        self.assertIn("推荐理由（活人感原句）", HEADERS)
        self.assertIn("原链接", HEADERS)

    def test_export_rows_exist_on_real_data(self):
        cards = collect_food_cards(latest_only=True)
        if not cards:  # 无真实数据时跳过（CI/干净机器）
            self.skipTest("无 food_cards 数据")
        from export_knowledge_base import collect_md_items

        md = collect_md_items()
        rows, stats = build_rows(cards, md, include_drinks=False)
        self.assertGreaterEqual(stats["rows"], 1)
        for row in rows:
            self.assertEqual(len(row), len(HEADERS))

    def test_workbook_writes_and_reopens(self):
        if not collect_food_cards(latest_only=True):
            self.skipTest("无 food_cards 数据")
        from export_knowledge_base import collect_md_items

        rows, stats = build_rows(collect_food_cards(latest_only=True), collect_md_items(), include_drinks=False)
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "kb.xlsx"
            write_xlsx(rows, out, stats, include_drinks=False)
            self.assertTrue(out.exists() and out.stat().st_size > 2000)
            try:
                import openpyxl

                wb = openpyxl.load_workbook(out)
                self.assertIn("新品知识库", wb.sheetnames)
                self.assertIn("口径说明", wb.sheetnames)
            except ImportError:
                pass


if __name__ == "__main__":
    unittest.main()
